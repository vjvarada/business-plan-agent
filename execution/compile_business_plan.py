#!/usr/bin/env python3
"""Compile a project's business plan sections into one document.

Outputs (local files):
  <out>/<Name>_Business_Plan.md    one markdown file, reference anchors made unique per section
  <out>/<Name>_Business_Plan.docx  Word version (python-docx; skipped if not installed)

Usage:
  python execution/compile_business_plan.py --project aroh --title "Aroh — Business Plan"

Section order is the filename order (00_..., 01_..., 11_...). Each section keeps its own
"References & Sources" list; anchors become s<NN>-ref-<N> so links never collide, and the
Word version renders citations as [N] within the section's own reference list.
"""
import argparse
import re
import sys
from datetime import date
from pathlib import Path

CITE = re.compile(r"\[\[(\d+)\]\]\(#ref-(\d+)\)")
ANCHOR = re.compile(r'<a name="ref-(\d+)"></a>')


def load_sections(sections_dir):
    files = sorted(p for p in Path(sections_dir).glob("[0-9][0-9]_*.md"))
    if not files:
        raise SystemExit(f"No numbered sections in {sections_dir}")
    return files


def section_title(text, fallback):
    m = re.search(r"^#\s+(.+)$", text, re.M)
    return m.group(1).strip() if m else fallback


def uniquify(text, num):
    text = ANCHOR.sub(lambda m: f'<a name="s{num}-ref-{m.group(1)}"></a>', text)
    return CITE.sub(lambda m: f"[[{m.group(1)}]](#s{num}-ref-{m.group(2)})", text)


def build_markdown(files, title, subtitle):
    parts = [f"# {title}\n", f"{subtitle}\n", f"> Compiled {date.today().isoformat()} from {len(files)} sections.\n", "## Contents\n"]
    toc, bodies = [], []
    for f in files:
        num = f.name[:2]
        text = f.read_text(encoding="utf-8")
        name = section_title(text, f.stem)
        toc.append(f"- [{num} — {name}](#section-{num})")
        body = uniquify(text, num)
        body = re.sub(r"^#\s+(.+)$", lambda m: f'<a name="section-{num}"></a>\n\n# {num} — {m.group(1)}', body, count=1, flags=re.M)
        bodies.append(body.strip() + "\n")
    return "\n".join(parts) + "\n".join(toc) + "\n\n---\n\n" + "\n\n---\n\n".join(bodies)


# ---------------------------------------------------------------- Word output
def add_runs(paragraph, text):
    """Add text with **bold** and *italic* spans; citations become [N]."""
    text = CITE.sub(lambda m: f"[{m.group(1)}]", text)
    text = ANCHOR.sub("", text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            paragraph.add_run(tok[1:-1]).italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            paragraph.add_run(tok[1:-1])
        else:
            paragraph.add_run(tok)


def build_docx(files, title, subtitle, out_path):
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        print("  python-docx not installed — skipping .docx")
        return False
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle)
    doc.add_paragraph(f"Compiled {date.today().isoformat()} from {len(files)} sections.")
    for f in files:
        num = f.name[:2]
        lines = f.read_text(encoding="utf-8").split("\n")
        doc.add_page_break()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if not line.strip() or ANCHOR.fullmatch(line.strip()) or line.strip() == "---":
                i += 1
                continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                if rows:
                    ncol = max(len(r) for r in rows)
                    table = doc.add_table(rows=len(rows), cols=ncol)
                    table.style = "Light Grid Accent 1"
                    for ri, r in enumerate(rows):
                        for ci in range(ncol):
                            cell = table.cell(ri, ci)
                            cell.text = ""
                            add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                            if ri == 0:
                                for run in cell.paragraphs[0].runs:
                                    run.bold = True
                    doc.add_paragraph()
                continue
            m = re.match(r"^(#{1,4})\s+(.*)$", line)
            if m:
                level = len(m.group(1))
                text = m.group(2)
                if level == 1:
                    text = f"{num} — {text}"
                doc.add_heading(text, min(level, 3))
            elif line.startswith("> "):
                para = doc.add_paragraph(style="Intense Quote")
                add_runs(para, line[2:])
            elif re.match(r"^\s*[-*]\s+", line):
                para = doc.add_paragraph(style="List Bullet")
                add_runs(para, re.sub(r"^\s*[-*]\s+", "", line))
            elif re.match(r"^\s*\d+\.\s+", line):
                para = doc.add_paragraph(style="List Number")
                add_runs(para, re.sub(r"^\s*\d+\.\s+", "", line))
            else:
                para = doc.add_paragraph()
                add_runs(para, line)
            i += 1
    doc.save(out_path)
    return True


def main():
    ap = argparse.ArgumentParser(description="Compile business plan sections into one document")
    ap.add_argument("--project", required=True, help="Project name under .tmp/")
    ap.add_argument("--title", help="Document title")
    ap.add_argument("--subtitle", default="", help="Subtitle / one-line description")
    ap.add_argument("--sections-dir")
    ap.add_argument("--out-dir")
    args = ap.parse_args()

    base = Path(".tmp") / args.project
    sections_dir = Path(args.sections_dir) if args.sections_dir else base / "business_plan" / "sections"
    out_dir = Path(args.out_dir) if args.out_dir else base / "business_plan"
    out_dir.mkdir(parents=True, exist_ok=True)
    name = args.project.capitalize()
    title = args.title or f"{name} — Business Plan"
    files = load_sections(sections_dir)

    md_path = out_dir / f"{name}_Business_Plan.md"
    md_path.write_text(build_markdown(files, title, args.subtitle), encoding="utf-8")
    print(f"✓ {md_path} ({len(files)} sections)")
    docx_path = out_dir / f"{name}_Business_Plan.docx"
    if build_docx(files, title, args.subtitle, docx_path):
        print(f"✓ {docx_path}")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
